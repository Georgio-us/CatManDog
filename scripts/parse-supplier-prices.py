"""Read saved supplier tariffs into private, reviewable JSON snapshots.

The output stays under the ignored research/supplier-files directory. No
supplier price is copied into a public sale price or confirmed purchase cost.
"""

import json
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "research/supplier-files"
OUTPUT = BASE / "parsed-price-snapshots"


def cents(value):
    if isinstance(value, str):
        value = value.replace("€", "").strip().replace(".", "").replace(",", ".")
    if isinstance(value, bool) or not isinstance(value, (int, float, str)) or value in ("", "-"):
        return None
    try:
        return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except Exception:
        return None


def text(value):
    return " ".join(str(value or "").split())


def ean(value):
    value = str(value or "").strip()
    return value if re.fullmatch(r"\d{8,14}", value) else ""


def snapshot(identifier, supplier, brand, date, file, shop_basis, items, note=""):
    articles = set()
    for item in items:
        if item["article"] in articles:
            raise ValueError(f"Duplicate article in {identifier}: {item['article']}")
        articles.add(item["article"])
        if not item["name"] or item["shopPriceCents"] is None and item["recommendedPriceCents"] is None:
            raise ValueError(f"Incomplete row in {identifier}: {item}")
    payload = dict(id=identifier, supplierName=supplier, brand=brand, priceDate=date,
                   sourceFile=file, shopPriceTaxBasis=shop_basis,
                   retailPriceTaxBasis="includes_iva", items=items)
    if note:
        payload["sourceNote"] = note
    OUTPUT.mkdir(parents=True, exist_ok=True)
    target = OUTPUT / f"{identifier}.json"
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n")
    print(f"{target.name}: {len(items)} rows, {sum(bool(item['ean']) for item in items)} EAN")


def royal_canin():
    file = BASE / "2026-10-14-ownat-royal-canin/TARIFAS TIENDA 2026 COSTE.xlsx"
    # The matching 01.03.2026 PDF tariff corrects three duplicated EAN in XLSX.
    corrected_eans = {"2554001": "3182550721202", "2554201": "3182550721219",
                      "2554401": "3182550721226"}
    sheet = load_workbook(file, read_only=True, data_only=True).active
    items = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_number < 4 or not re.fullmatch(r"[A-Z0-9]{4,20}", str(row[4] or "")) or not isinstance(row[8], (int, float)):
            continue
        species, technology = text(row[1]), text(row[2])
        unit = text(row[7])
        article = str(row[4])
        items.append(dict(article=article, name=text(row[5]),
                          line=f"{species} / {text(row[0])} / {technology}",
                          weight=f"{row[6]} {unit}" if row[6] is not None else "",
                          ean=corrected_eans.get(article, ean(row[3])), shopPriceCents=cents(row[8]),
                          recommendedPriceCents=None, minimumPriceCents=None,
                          rawPrices=dict(kind="supplier_tariff_ex_iva", sourceAmount=row[8],
                                         unit="as_listed", species=species,
                                         eanCorrectionSource="Tarifas-PSR-2026.pdf" if article in corrected_eans else None,
                                         note="Check case versus individual unit before using as purchase cost"),
                          sourceRow=row_number))
    snapshot("royal-canin-2026-03-01", "Royal Canin", "Royal Canin", "2026-03-01",
             file.name, "excludes_iva", items)


def belcando():
    file = BASE / "2026-10-28-leonardo/Belcando  01.03.2026.xlsx"
    sheet = load_workbook(file, read_only=True, data_only=True).active
    items, line = [], ""
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row[0] is None and isinstance(row[1], str) and "Belcando" in row[1]:
            line = text(row[1])
        if not isinstance(row[0], int) or not isinstance(row[6], (int, float)):
            continue
        items.append(dict(article=str(row[0]), name=text(row[1]), line=line,
                          weight=text(row[2]), ean=ean(row[3]),
                          shopPriceCents=cents(row[6]),
                          recommendedPriceCents=cents(row[7]),
                          minimumPriceCents=cents(row[8]),
                          rawPrices=dict(kind="supplier_sheet_tienda", shop=row[6],
                                         recommended=row[7], minimum=row[8],
                                         unitsPerReference=row[5], caseEan=ean(row[4])),
                          sourceRow=row_number))
    snapshot("albet-belcando-2026-03-01", "Albet", "Belcando", "2026-03-01",
             file.name, "unknown", items)


def edgard_cooper():
    file = BASE / "2026-10-28-edgard-cooper/0. Pricelist_Spain 2026_NUEVA (2).xlsx"
    sheet = load_workbook(file, read_only=True, data_only=True).active
    items, line, name, species = [], "", "", "dog"
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        first = text(row[0])
        if "GATO" in first.upper() or "CAT" in first.upper():
            species = "cat"
        if row[2] is None and first and row_number > 8:
            line = first
        product_ean = ean(row[2])
        if not product_ean or cents(row[9]) is None and cents(row[10]) is None:
            continue
        if text(row[4]):
            name = text(row[4])
        if not name:
            raise ValueError(f"Missing name before row {row_number} in {file.name}")
        items.append(dict(article=product_ean, name=name, line=f"{species} / {line} / {text(row[6])}",
                          weight=text(row[5]), ean=product_ean,
                          shopPriceCents=cents(row[9]), recommendedPriceCents=cents(row[10]),
                          minimumPriceCents=None,
                          rawPrices=dict(kind="purchase_price_per_product_ex_iva",
                                         purchaseBoxCents=cents(row[8]),
                                         unitsPerBox=row[7], boxEan=ean(row[3]),
                                         recommendedExIvaCents=cents(row[11]),
                                         note="2026 price list; exact effective date not specified"),
                          sourceRow=row_number))
    snapshot("edgard-cooper-2026-01-01", "Edgard & Cooper", "Edgard & Cooper",
             "2026-01-01", file.name, "excludes_iva", items,
             "2026 price list, exact effective date not specified")


def ownat():
    folder = BASE / "2026-10-14-ownat-royal-canin"
    files = sorted(list(folder.glob("ON LINE *2026.pdf")) + list(folder.glob("ONLINE *2026*.pdf")))
    for file in files:
        slug = re.sub(r"[^a-z0-9]+", "-", file.stem.lower()).strip("-")
        items = []
        last_name = ""
        line = file.stem
        row_number = 0
        for page in pdfplumber.open(file).pages:
            for table in page.extract_tables():
                for row in table:
                    row_number += 1
                    cells = [text(cell) for cell in row]
                    code_index = next((i for i, cell in enumerate(cells) if re.fullmatch(r"\d{6}", cell)), None)
                    if code_index is None or len(cells) < code_index + 6:
                        continue
                    code = cells[code_index]
                    product = cells[code_index + 1]
                    product_ean = ean(cells[code_index + 2])
                    if not product_ean:
                        raise ValueError(f"Missing EAN in {file.name} row {row_number}: {cells}")
                    if product:
                        last_name = product
                    if not last_name:
                        raise ValueError(f"Missing product before {code} in {file.name}")
                    wet = "WETLINE" in file.name.upper()
                    price_index = code_index + (5 if wet else 4)
                    if len(cells) <= price_index + 1:
                        raise ValueError(f"Missing prices for {code} in {file.name}")
                    shop = cents(cells[price_index])
                    retail = cents(cells[price_index + 1])
                    if shop is None or retail is None:
                        raise ValueError(f"Bad prices for {code} in {file.name}: {cells}")
                    units = int(cells[code_index + 4]) if wet else 1
                    items.append(dict(article=code, name=last_name, line=line,
                                      weight=f"{cells[code_index + 3]} {'g' if wet else 'kg'}",
                                      ean=product_ean,
                                      shopPriceCents=None if wet else shop,
                                      recommendedPriceCents=retail, minimumPriceCents=None,
                                      rawPrices=dict(kind="recommended_shop_ex_iva",
                                                     recommendedShopCaseCents=shop if wet else None,
                                                     unitsPerCase=units, sourceShopCents=shop,
                                                     note="Recommended shop price, not confirmed CatManDog purchase cost"),
                                      sourceRow=row_number))
        snapshot(f"ownat-2026-02-01-{slug}", "Comercial Química Muñoz", "Ownat",
                 "2026-02-01", file.name, "excludes_iva", items)


if __name__ == "__main__":
    royal_canin()
    belcando()
    edgard_cooper()
    ownat()
