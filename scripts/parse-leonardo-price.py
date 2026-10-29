"""Extract the Albet Leonardo price list into an internal, versioned JSON snapshot.

Run with a Python environment that has openpyxl installed. This script never
changes the supplier workbook or the public shop catalog.
"""

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "research/supplier-files/2026-10-28-leonardo/Copia de Leonardo 01.03.2026.xlsx"
OUTPUT = SOURCE.with_name("albet-leonardo-2026-03-01.json")


def cents(value):
    if not isinstance(value, (float, int)):
        return None
    return int((Decimal(str(value)) * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def main():
    sheet = load_workbook(SOURCE, read_only=True, data_only=True).active
    line = ""
    items = []
    articles = set()
    for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
        article, name, weight, ean, shop_price = cells[:5]
        recommended, minimum = cells[11:13]
        if article is None and isinstance(name, str) and name.strip().lower().startswith("leonardo"):
            line = " ".join(name.split())
            continue
        if not isinstance(article, int) or not isinstance(name, str) or not isinstance(shop_price, (float, int)):
            continue
        article = str(article)
        if article in articles:
            raise ValueError(f"Duplicate supplier article: {article}")
        articles.add(article)
        items.append({
            "article": article,
            "name": " ".join(name.split()),
            "line": line,
            "weight": str(weight or "").strip(),
            "ean": str(ean).strip() if ean is not None else "",
            "shopPriceCents": cents(shop_price),
            "recommendedPriceCents": cents(recommended),
            "minimumPriceCents": cents(minimum),
            "rawPrices": {"shop": shop_price, "recommended": recommended, "minimum": minimum},
            "sourceRow": row_number,
        })
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps({
        "id": "albet-leonardo-2026-03-01",
        "supplierName": "Albet",
        "brand": "Leonardo",
        "priceDate": "2026-03-01",
        "sourceFile": SOURCE.name,
        "shopPriceTaxBasis": "unknown",
        "retailPriceTaxBasis": "includes_iva",
        "items": items,
    }, ensure_ascii=False, indent=2) + "\n")
    print(f"Saved {len(items)} supplier rows to {OUTPUT}")


if __name__ == "__main__":
    main()
